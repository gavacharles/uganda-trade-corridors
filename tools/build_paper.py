"""Build a Word paper from a small Markdown subset.

    python tools/build_paper.py paper/paper.md paper/uganda_corridors_paper.docx

Front matter (top of the file, one per line): `title:`, `subtitle:`, `author:`, `affil:`, `email:`,
`keywords:`. Then:
  # Heading / ## Subheading / ### Minor heading
  plain lines (joined until a blank line) = a paragraph; **bold**, *italic* inline
  - item               bullet
  ![Figure 1. Caption](relative/path.png){6.3}    figure, width in inches (default 6.3)
  Table: Table 1. Caption    then a pipe table (| a | b |, second row |---|)
  <<pagebreak>>
Under a heading named "References", paragraphs get a hanging indent. Figure and table paths are
relative to the Markdown file.
"""
import os, re, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FONT, BODY_PT = "Times New Roman", 11
INK = RGBColor(0x1D, 0x23, 0x21)
GREY = RGBColor(0x5D, 0x67, 0x64)


def setup(doc):
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Cm(2.5))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_after, pf.line_spacing = Pt(6), 1.2
    for name, size, before in (("Heading 1", 14, 18), ("Heading 2", 12, 12), ("Heading 3", 11, 10)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold, h.font.color.rgb = FONT, Pt(size), True, INK
        h.font.italic = name == "Heading 3"
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(4)
        h.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    cap.font.name, cap.font.size, cap.font.italic, cap.font.bold = FONT, Pt(9.5), False, False
    cap.font.color.rgb = INK
    cap.paragraph_format.space_after = Pt(10)
    # page numbers, centred in the footer
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)
    r.font.size = Pt(9)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def runs(p, text, size=None, color=None, bold=None, italic=None):
    for part in INLINE.split(text):
        if not part:
            continue
        b = part.startswith("**")
        i = not b and part.startswith("*") and part.endswith("*")
        r = p.add_run(part[2:-2] if b else part[1:-1] if i else part)
        r.bold = True if b else bold
        r.italic = True if i else italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_)
    tcPr.append(sh)


def borders(table):
    """Three-line academic table: rules above and below the header and at the bottom."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge, val in (("top", "single"), ("bottom", "single"), ("left", "nil"), ("right", "nil"),
                      ("insideH", "nil"), ("insideV", "nil")):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), "8")
        e.set(qn("w:color"), "1D2321")
        b.append(e)
    tblPr.append(b)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tb = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "1D2321")
        tb.append(e)
        tcPr.append(tb)


def add_table(doc, caption, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    cp = doc.add_paragraph(style="Caption")
    runs(cp, caption, size=9.5)
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.space_after = Pt(3)
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            numeric = bool(re.fullmatch(r"[−\-–+~≈<>$]?[\d.,]+[%×]?(\s*[–-]\s*[\d.,]+%?)?( ?[a-zA-Z%]{0,4})?", c))
            if j > 0 and numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            runs(p, c, size=9, bold=True if i == 0 else None)
            if i == 0:
                shade(cell, "EFEDE8")
    borders(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build(md_path, out):
    base = os.path.dirname(os.path.abspath(md_path))
    text = open(md_path, encoding="utf-8").read()
    # Figure labels (@name) are numbered in the order their images appear, so inserting a figure
    # renumbers every reference to the ones after it
    order = re.findall(r"^!\[Figure @(\w+)\.", text, flags=re.M)
    num = {k: str(i + 1) for i, k in enumerate(order)}
    missing = set(re.findall(r"@(\w+)", text)) - set(num) - {"gmail"}
    if missing:
        sys.exit(f"figure labels referenced but not placed: {sorted(missing)}")
    text = re.sub(r"@(\w+)", lambda m: num.get(m.group(1), m.group(0)), text)
    lines = text.split("\n")
    meta = {}
    while lines and re.match(r"^(title|subtitle|author|affil|email|keywords):", lines[0]):
        k, v = lines.pop(0).split(":", 1)
        meta[k] = v.strip()
    doc = Document()
    setup(doc)
    # title block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(p, meta.get("title", ""), size=17, bold=True, color=INK)
    if "subtitle" in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(p, meta["subtitle"], size=12.5, italic=True, color=GREY)
    for k, size in (("author", 11.5), ("affil", 10), ("email", 10)):
        if k in meta:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0 if k != "email" else 12)
            runs(p, meta[k], size=size, color=INK if k == "author" else GREY)
    in_refs, para, i = False, [], 0

    def flush():
        nonlocal para
        if para:
            text = " ".join(x.strip() for x in para)
            p = doc.add_paragraph()
            if in_refs:
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.first_line_indent = Cm(-0.8)
                p.paragraph_format.space_after = Pt(3)
                runs(p, text, size=10)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                runs(p, text)
            para = []

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush()
        elif s == "<<pagebreak>>":
            flush()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            text = s[level:].strip()
            doc.add_heading(text, level=level)
            in_refs = text.lower().startswith("references")
        elif s.startswith("- "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            runs(p, s[2:])
        elif s.startswith("!["):
            flush()
            m = re.match(r"!\[(.*)\]\((.+?)\)(\{([\d.]+)\})?", s)
            cap, path, w = m.group(1), m.group(2), float(m.group(4) or 6.3)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(os.path.join(base, path), width=Inches(w))
            cp = doc.add_paragraph(style="Caption")
            runs(cp, cap, size=9.5)
        elif s.startswith("Table:"):
            flush()
            cap = s[len("Table:"):].strip()
            rows = []
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, cap, rows)
            continue
        elif s.startswith("Keywords:") or s.startswith("**Keywords"):
            flush()
            p = doc.add_paragraph()
            runs(p, s)
        else:
            para.append(ln)
        i += 1
    flush()
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
